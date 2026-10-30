"""
Enterprise Excel storage and styling manager utilizing Pandas and Openpyxl.
Creates two workbooks:
1. OrangeHRM_Data.xlsx (Timesheet Data)
2. Details.xlsx (10 My Info Worksheets, each containing its section data + Attachments table with PDF hyperlinks)
"""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.exceptions import ExcelExportError

logger = logging.getLogger(__name__)


class ExcelManager:
    """Manages writing and professionally formatting Excel workbooks."""

    HEADER_FILL = PatternFill(
        start_color="1F4E78", end_color="1F4E78", fill_type="solid"
    )
    SECTION_BANNER_FILL = PatternFill(
        start_color="D9E1F2", end_color="D9E1F2", fill_type="solid"
    )
    ZEBRA_FILL = PatternFill(
        start_color="F9FAFB", end_color="F9FAFB", fill_type="solid"
    )

    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    SECTION_FONT = Font(name="Calibri", size=11, bold=True, color="1F4E78")
    DATA_FONT = Font(name="Calibri", size=10, bold=False, color="000000")
    LINK_FONT = Font(
        name="Calibri", size=10, bold=True, color="0563C1", underline="single"
    )

    THIN_BORDER = Border(
        left=Side(style="thin", color="D3D3D3"),
        right=Side(style="thin", color="D3D3D3"),
        top=Side(style="thin", color="D3D3D3"),
        bottom=Side(style="thin", color="D3D3D3"),
    )

    @classmethod
    def save_to_sheet(
        cls, file_path: Path, sheet_name: str, data: list[dict[str, Any]]
    ) -> None:
        """
        Write records to an Excel sheet and apply table styling.

        Args:
            file_path (Path): Destination workbook path.
            sheet_name (str): Target worksheet name.
            data (List[Dict[str, Any]]): Rows to write.

        Returns:
            None.

        Example:
            ExcelManager.save_to_sheet(path, "Timesheet_Data", rows)
        """
        if not data:
            data = [{"Status": "No records found in UI table"}]

        clean_sheet_name = sheet_name[:31]
        df = pd.DataFrame(data)

        try:
            if file_path.exists():
                with pd.ExcelWriter(
                    file_path,
                    engine="openpyxl",
                    mode="a",
                    if_sheet_exists="replace",
                ) as writer:
                    df.to_excel(writer, sheet_name=clean_sheet_name, index=False)
            else:
                with pd.ExcelWriter(file_path, engine="openpyxl", mode="w") as writer:
                    df.to_excel(writer, sheet_name=clean_sheet_name, index=False)

            cls._apply_professional_styling(file_path, clean_sheet_name)
            logger.info(
                "Saved and styled %d row(s) in '%s' -> [%s]",
                len(df),
                file_path.name,
                clean_sheet_name,
            )
        except (OSError, KeyError, TypeError, ValueError) as exc:
            logger.error(
                "Failed to write sheet '%s' in '%s': %s",
                clean_sheet_name,
                file_path,
                exc,
            )
            raise ExcelExportError(
                f"Failed to export sheet '{clean_sheet_name}' to '{file_path.name}': {exc}"
            ) from exc

    @classmethod
    def save_section_with_attachments(
        cls,
        file_path: Path,
        sheet_name: str,
        section_data: list[dict[str, Any]],
        attachments_data: list[dict[str, Any]],
        pdf_path: Path | None = None,
    ) -> None:
        """
        Write section data and an attachments block to one worksheet.

        Args:
            file_path (Path): Destination workbook path.
            sheet_name (str): Target worksheet name.
            section_data (List[Dict[str, Any]]): Primary section rows.
            attachments_data (List[Dict[str, Any]]): Attachment metadata rows.
            pdf_path (Optional[Path]): Local PDF target for hyperlinks.

        Returns:
            None.

        Example:
            ExcelManager.save_section_with_attachments(path, "Job", rows, attachments)
        """
        if not section_data:
            section_data = [{"Status": "No records found in UI section"}]

        clean_sheet_name = sheet_name[:31]

        try:
            if file_path.exists():
                wb = load_workbook(file_path)
                if clean_sheet_name in wb.sheetnames:
                    del wb[clean_sheet_name]
                ws = wb.create_sheet(title=clean_sheet_name)
            else:
                wb = Workbook()
                ws = wb.active
                ws.title = clean_sheet_name

            # -----------------------------------------------------------------
            # PART 1: Write Primary Section Data Table
            # -----------------------------------------------------------------
            section_headers = list(section_data[0].keys())
            current_row = 1

            # Section Data Banner
            max_cols = max(len(section_headers), 6)
            ws.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=max_cols,
            )
            sec_banner = ws.cell(
                row=current_row,
                column=1,
                value=f"MY INFO SECTION: {clean_sheet_name.replace('_', ' ').upper()}",
            )
            sec_banner.font = cls.SECTION_FONT
            sec_banner.fill = cls.SECTION_BANNER_FILL
            sec_banner.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[current_row].height = 24
            for col_idx in range(1, max_cols + 1):
                ws.cell(row=current_row, column=col_idx).border = cls.THIN_BORDER
            current_row += 1

            # Primary Data Column Headers
            ws.row_dimensions[current_row].height = 22
            for col_idx, h_text in enumerate(section_headers, start=1):
                c = ws.cell(row=current_row, column=col_idx, value=h_text)
                c.fill = cls.HEADER_FILL
                c.font = cls.HEADER_FONT
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border = cls.THIN_BORDER
            current_row += 1

            # Primary Data Rows
            for row_idx, record in enumerate(section_data):
                ws.row_dimensions[current_row].height = 20
                is_even = row_idx % 2 == 1
                for col_idx, h_text in enumerate(section_headers, start=1):
                    val = record.get(h_text, "N/A")
                    c = ws.cell(row=current_row, column=col_idx, value=val)
                    c.font = cls.DATA_FONT
                    c.border = cls.THIN_BORDER
                    if is_even:
                        c.fill = cls.ZEBRA_FILL
                    c.alignment = Alignment(horizontal="left", vertical="center")
                current_row += 1

            # Leave 2 blank spacing rows before Attachments block
            current_row += 2

            # -----------------------------------------------------------------
            # PART 2: Write Section Attachments Block Under Heading "ATTACHMENTS"
            # -----------------------------------------------------------------
            att_headers = [
                "File Name",
                "Description",
                "Size",
                "Type",
                "Date Added",
                "Added By",
            ]

            # Attachments Heading Banner
            ws.merge_cells(
                start_row=current_row,
                start_column=1,
                end_row=current_row,
                end_column=len(att_headers),
            )
            att_banner = ws.cell(
                row=current_row,
                column=1,
                value=f"ATTACHMENTS ({clean_sheet_name.replace('_', ' ').upper()})",
            )
            att_banner.font = cls.SECTION_FONT
            att_banner.fill = cls.SECTION_BANNER_FILL
            att_banner.alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[current_row].height = 24
            for col_idx in range(1, len(att_headers) + 1):
                ws.cell(row=current_row, column=col_idx).border = cls.THIN_BORDER
            current_row += 1

            # Attachments Column Headers
            ws.row_dimensions[current_row].height = 22
            for col_idx, h_text in enumerate(att_headers, start=1):
                c = ws.cell(row=current_row, column=col_idx, value=h_text)
                c.fill = cls.HEADER_FILL
                c.font = cls.HEADER_FONT
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border = cls.THIN_BORDER
            current_row += 1

            # Resolve local PDF path for clickable Hyperlink
            local_pdf_target = (
                str(pdf_path.resolve())
                if pdf_path
                else str(Path("devdata/sample_attachment.pdf").resolve())
            )

            # Attachments Data Rows
            if attachments_data:
                for att_idx, rec in enumerate(attachments_data):
                    ws.row_dimensions[current_row].height = 20
                    is_even = att_idx % 2 == 1
                    row_vals = [
                        rec.get("File Name", "N/A"),
                        rec.get("Description", "N/A"),
                        rec.get("Size", "N/A"),
                        rec.get("Type", "N/A"),
                        rec.get("Date Added", "N/A"),
                        rec.get("Added By", "N/A"),
                    ]
                    for col_idx, val in enumerate(row_vals, start=1):
                        c = ws.cell(row=current_row, column=col_idx, value=val)
                        c.border = cls.THIN_BORDER
                        if is_even:
                            c.fill = cls.ZEBRA_FILL
                        c.alignment = Alignment(
                            horizontal="left" if col_idx <= 2 else "center",
                            vertical="center",
                        )

                        # Make 'File Name' (Column 1) a clickable hyperlink to the PDF
                        if col_idx == 1:
                            c.hyperlink = local_pdf_target
                            c.font = cls.LINK_FONT
                        else:
                            c.font = cls.DATA_FONT
                    current_row += 1
            else:
                ws.row_dimensions[current_row].height = 20
                c = ws.cell(
                    row=current_row,
                    column=1,
                    value="No attachment widget / records in this section",
                )
                c.font = Font(name="Calibri", size=10, italic=True, color="666666")
                for col_idx in range(1, len(att_headers) + 1):
                    ws.cell(row=current_row, column=col_idx).border = cls.THIN_BORDER
                current_row += 1

            cls._auto_fit_columns(ws)
            wb.save(file_path)
            logger.info(
                "Saved section [%s] + Attachments with PDF hyperlink into '%s'",
                clean_sheet_name,
                file_path.name,
            )

        except (OSError, KeyError, TypeError, ValueError) as exc:
            logger.error(
                "Failed to save section '%s' with attachments in '%s': %s",
                clean_sheet_name,
                file_path,
                exc,
            )
            raise ExcelExportError(
                f"Failed to write section '{clean_sheet_name}' in '{file_path.name}': {exc}"
            ) from exc

    @classmethod
    def _apply_professional_styling(cls, file_path: Path, sheet_name: str) -> None:
        """Apply workbook styling to a worksheet.

        Args:
            file_path (Path): Workbook path.
            sheet_name (str): Worksheet to style.

        Returns:
            None.

        Example:
            ExcelManager._apply_professional_styling(path, "Timesheet_Data")
        """
        wb = load_workbook(file_path)
        if sheet_name not in wb.sheetnames:
            return
        ws = wb[sheet_name]

        # Style Header Row (Row 1)
        ws.row_dimensions[1].height = 24
        for cell in ws[1]:
            cell.fill = cls.HEADER_FILL
            cell.font = cls.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = cls.THIN_BORDER

        # Style Data Rows
        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 20
            is_even = row_idx % 2 == 0
            for cell in ws[row_idx]:
                cell.font = cls.DATA_FONT
                cell.border = cls.THIN_BORDER
                if is_even:
                    cell.fill = cls.ZEBRA_FILL
                cell.alignment = Alignment(horizontal="left", vertical="center")

        cls._auto_fit_columns(ws)
        wb.save(file_path)

    @staticmethod
    def _auto_fit_columns(ws: Any) -> None:
        """Adjust worksheet columns to fit their longest text.

        Args:
            ws (Any): Openpyxl worksheet object.

        Returns:
            None.

        Example:
            ExcelManager._auto_fit_columns(worksheet)
        """
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or "")
                # Skip merged banner headers when calculating width
                if val_str.startswith("MY INFO SECTION:") or val_str.startswith(
                    "ATTACHMENTS ("
                ):
                    continue
                max_len = max(max_len, len(val_str))
            ws.column_dimensions[col_letter].width = max(max_len + 5, 15)
